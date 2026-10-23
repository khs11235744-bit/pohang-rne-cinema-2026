# -*- coding: utf-8 -*-
"""Export four printable A4 student speech scripts with Q&A to DOCX."""
from pathlib import Path
import json, math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path.home()/'Documents'/'KHS_RNE_2026'/'MultiTeamCinema_20261023'/'v2-studio'
OUT=root/'public'/'docs';OUT.mkdir(parents=True,exist_ok=True)
DOCS=root/'docs';DOCS.mkdir(parents=True,exist_ok=True)
teams=['shin','cha','hu','seung']
opts={
 'shin':('신진욱팀','1인 기본 / 2인 선택',RGBColor(25,112,89)),
 'cha':('차민혁팀','1인 기본 / 2인 선택',RGBColor(156,40,124)),
 'hu':('허준완팀','허준완·김주찬 2인 기본',RGBColor(141,85,37)),
 'seung':('신승연팀','1인 기본 / 2인 선택',RGBColor(22,125,120))
}
def east(run,font='맑은 고딕'):
 run.font.name=font
 rPr=run._element.get_or_add_rPr()
 rFonts=rPr.rFonts
 if rFonts is None:
  rFonts=OxmlElement('w:rFonts');rPr.insert(0,rFonts)
 rFonts.set(qn('w:eastAsia'),font)
def add_text(p,text,size=None,bold=False,color=None):
 r=p.add_run(str(text))
 east(r)
 if size:r.font.size=Pt(size)
 r.bold=bold
 if color:r.font.color.rgb=color
 return r
def add_page_number(para):
 para.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 add_text(para,'2026 POHANG R&E  |  ',8)
 r=para.add_run()
 fld=OxmlElement('w:fldSimple')
 fld.set(qn('w:instr'),'PAGE')
 r._r.addnext(fld)
def keep_with_next(p):
 p.paragraph_format.keep_with_next=True
def set_border_bottom(p,color='CBCFCC'):
 pPr=p._element.get_or_add_pPr()
 borders=OxmlElement('w:pBdr');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'12');b.set(qn('w:color'),color);borders.append(b);pPr.append(borders)
for id in teams:
 data=json.loads((root/'data'/(id+'.json')).read_text('utf8'))
 name,mode,accent=opts[id]
 doc=Document()
 sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
 sec.header_distance=Cm(.75);sec.footer_distance=Cm(.75)
 st=doc.styles['Normal'];st.font.name='맑은 고딕';st.font.size=Pt(10);st.font.color.rgb=RGBColor(36,55,65)
 st.paragraph_format.space_after=Pt(7)
 for h,sz in [('Title',22),('Heading 1',15),('Heading 2',12)]:
  stt=doc.styles[h];stt.font.name='맑은 고딕';stt.font.size=Pt(sz);stt.font.bold=True;stt.font.color.rgb=accent
  stt.paragraph_format.space_before=Pt(16);stt.paragraph_format.space_after=Pt(6)
 header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 add_text(header,'2026 포항고등학교 · 인문사회 R&E',8,color=accent)
 add_page_number(sec.footer.paragraphs[0])
 p=doc.add_paragraph();add_text(p,'2026 포항고 인문사회 R&E  |  학생 발표 대본',10,bold=True,color=accent)
 p=doc.add_paragraph(style='Title');add_text(p,name+' · 7분 발표 대본',22,bold=True,color=accent)
 p=doc.add_paragraph();add_text(p,data['title'],15,bold=True)
 p=doc.add_paragraph();add_text(p,'연구 주제  ',10,bold=True,color=accent);add_text(p,data['project'],10)
 p=doc.add_paragraph();add_text(p,'참여학생  ',10,bold=True,color=accent);add_text(p,' · '.join(data['students']),10)
 p=doc.add_paragraph();add_text(p,'발표 인원  ',10,bold=True,color=accent);add_text(p,mode,10)
 p=doc.add_paragraph();add_text(p,f"슬라이드  {len(data['slides'])}장     대본 권장시간  420초 (7분)     예상 질문  {len(data['questions'])}개",10,bold=True)
 p=doc.add_paragraph();add_text(p,'주의: 아래 장별 시간은 목표 시간입니다. 학생의 실제 말하기 속도와 화면 효과 재생시간은 리허설에서 확인해야 합니다.',9,color=RGBColor(114,99,76))
 set_border_bottom(p)
 doc.add_paragraph('슬라이드별 발표 대본',style='Heading 1')
 for i,s in enumerate(data['slides']):
  p=doc.add_paragraph(style='Heading 2');keep_with_next(p)
  add_text(p,f"{i+1:02d}. {s['title']}  |  {s['duration']}초",12,bold=True,color=accent)
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6)
  add_text(p,'1인: 대표 발표자  ·  2인: '+s['speakerTwo'],9,bold=True,color=RGBColor(90,112,119))
  p=doc.add_paragraph()
  add_text(p,s['notes'],10)
  p=doc.add_paragraph();add_text(p,'화면 주장  ',9,bold=True,color=accent);add_text(p,s['claim'],9)
  p=doc.add_paragraph();add_text(p,'자료 근거  ',8,bold=True,color=RGBColor(95,110,117));add_text(p,s.get('source',''),8,color=RGBColor(95,110,117))
 doc.add_page_break()
 doc.add_paragraph('심사위원 예상 질문·모범답안',style='Heading 1')
 p=doc.add_paragraph();add_text(p,'이 문항은 실전 연습용입니다. 설문 문항의 분모, 학생의 인식과 실제 효과, 출처의 한계를 함께 설명하도록 연습하세요.',10)
 for i,q in enumerate(data['questions']):
  p=doc.add_paragraph(style='Heading 2')
  add_text(p,f"{i+1:02d}. {q['question']}",11,bold=True,color=accent)
  keep_with_next(p)
  p=doc.add_paragraph();add_text(p,'모범답안  ',9,bold=True,color=accent);add_text(p,q['answer'],10)
 doc.add_paragraph('연구자료·학자 출처',style='Heading 1')
 for source in data['sources']:
  p=doc.add_paragraph(style='Normal')
  add_text(p,source['name']+' — ',9,bold=True)
  add_text(p,source['detail'],9)
  if source.get('url'):
   p=doc.add_paragraph(style='Normal');add_text(p,source['url'],8,color=RGBColor(80,119,120))
 p=doc.add_paragraph()
 add_text(p,'연구와 발표의 한계: 한 학교의 설문은 전체 청소년의 대표성이 없으며, 학생 인식으로 인과관계를 입증할 수 없습니다. 본 발표의 95점은 제작 목표일 뿐 실제 심사 점수가 아닙니다.',8,color=RGBColor(97,92,87))
 file=OUT/(id+'_7min_RNE_script.docx')
 doc.save(file)
 assert file.exists() and file.stat().st_size>20000
 from zipfile import ZipFile
 with ZipFile(file) as z:assert z.testzip() is None
 print(id,'docx',file.stat().st_size,'bytes','slides',len(data['slides']),'qa',len(data['questions']))
